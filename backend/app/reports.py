from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Flowable, Image as PDFImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class BrandMark(Flowable):
    def __init__(self) -> None:
        super().__init__()
        self.width = 42
        self.height = 42

    def draw(self) -> None:
        self.canv.setFillColor(colors.HexColor("#ad6244"))
        self.canv.roundRect(0, 0, 38, 38, 10, fill=1, stroke=0)
        self.canv.setStrokeColor(colors.white)
        self.canv.setLineWidth(2.5)
        self.canv.setLineCap(1)
        self.canv.line(8, 28, 29, 28)
        self.canv.line(8, 20, 22, 20)
        self.canv.line(8, 12, 17, 12)
        self.canv.line(21, 12, 25, 8)
        self.canv.line(25, 8, 31, 17)


@dataclass
class QuoteDocument:
    job_title: str
    client_name: str
    trade: str
    option_name: str
    description: str
    price_cop: int
    duration: int
    duration_unit: str
    logo_path: str = ""
    logo_bytes: bytes = b""
    business_name: str = "Cotización de servicio"


def safe_filename(value: str) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9_-]+", "-", value).strip("-").lower()
    return (normalized[:50] or "cotizacion")


def _plain_text(data: QuoteDocument) -> bytes:
    lines = [
        "COTIZACIÓN — COTIZAFÁCIL",
        "",
        f"Trabajo: {data.job_title}",
        f"Cliente: {data.client_name}",
        f"Oficio: {data.trade}",
        "",
        f"Alternativa: {data.option_name}",
        data.description,
        "",
        f"Valor total (COP): {data.price_cop:,}".replace(",", "."),
        f"Duración indicada por el técnico: {data.duration} {data.duration_unit}",
        "",
        "La propuesta debe ser revisada y aprobada por el técnico antes de aceptarse.",
    ]
    return "\n".join(lines).encode("utf-8")


def _docx(data: QuoteDocument) -> bytes:
    document = Document()
    if data.logo_bytes:
        document.add_picture(io.BytesIO(data.logo_bytes), width=Inches(0.9))
    elif data.logo_path and Path(data.logo_path).is_file():
        document.add_picture(data.logo_path, width=Inches(0.9))
    document.add_paragraph(f"{data.business_name} · Hecho con CotizaFácil")
    document.add_heading("Cotización", 0)
    document.add_heading(data.job_title, level=1)
    document.add_paragraph(f"Preparada para: {data.client_name}")
    document.add_paragraph(f"Especialidad: {data.trade}")
    document.add_heading(data.option_name, level=2)
    document.add_paragraph(data.description)
    document.add_paragraph(f"Valor total: ${data.price_cop:,.0f} COP".replace(",", "."))
    document.add_paragraph(f"Duración estimada indicada por el técnico: {data.duration} {data.duration_unit}")
    document.add_paragraph("Esta propuesta debe ser revisada y aprobada por el técnico antes de aceptarse.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx(data: QuoteDocument) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cotización"
    rows = [
        [data.business_name, "Hecho con CotizaFácil"],
        ["Trabajo", data.job_title],
        ["Cliente", data.client_name],
        ["Oficio", data.trade],
        ["Alternativa", data.option_name],
        ["Alcance", data.description],
        ["Valor total (COP)", data.price_cop],
        ["Duración indicada por el técnico", f"{data.duration} {data.duration_unit}"],
        ["Nota", "Propuesta sujeta a revisión y aprobación del técnico."],
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 68
    sheet.column_dimensions["C"].width = 18
    sheet.row_dimensions[1].height = 42
    sheet["B7"].number_format = '"$" #,##0 "COP"'
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="556D48")
    if data.logo_bytes or (data.logo_path and Path(data.logo_path).is_file()):
        logo = ExcelImage(io.BytesIO(data.logo_bytes)) if data.logo_bytes else ExcelImage(data.logo_path)
        logo.width = 36
        logo.height = 36
        sheet.add_image(logo, "C1")
    for row in sheet.iter_rows(min_row=2, max_col=2):
        row[0].font = row[0].font.copy(bold=True, color="526B45")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[6].height = 56
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf(data: QuoteDocument) -> bytes:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    brand = ParagraphStyle("Brand", parent=styles["Heading2"], textColor=colors.HexColor("#3d5338"), fontSize=17, leading=20, spaceAfter=2)
    subtitle = ParagraphStyle("Subtitle", parent=styles["Normal"], textColor=colors.HexColor("#766b60"), fontSize=9)
    if data.logo_bytes:
        mark = PDFImage(io.BytesIO(data.logo_bytes), width=0.55 * inch, height=0.55 * inch, kind="proportional")
    elif data.logo_path and Path(data.logo_path).is_file():
        mark = PDFImage(data.logo_path, width=0.55 * inch, height=0.55 * inch, kind="proportional")
    else:
        mark = BrandMark()
    header = Table([[mark, [Paragraph(escape(data.business_name), brand), Paragraph("COTIZACIÓN DE SERVICIO · COTIZAFÁCIL", subtitle)]]], colWidths=[0.6 * inch, 5.7 * inch])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story = [
        header,
        Paragraph(f"Fecha de emisión: {date.today().strftime('%d/%m/%Y')}", subtitle),
        Spacer(1, 0.25 * inch),
        Paragraph(escape(data.job_title), styles["Heading1"]),
    ]
    table = Table([
        ["Cliente", Paragraph(escape(data.client_name), styles["BodyText"])],
        ["Especialidad", Paragraph(escape(data.trade), styles["BodyText"])],
        ["Alternativa", Paragraph(escape(data.option_name), styles["BodyText"])],
        ["Alcance", Paragraph(escape(data.description).replace("\n", "<br/>"), styles["BodyText"])],
        ["Valor total", f"${data.price_cop:,.0f} COP".replace(",", ".")],
        ["Duración estimada", f"{data.duration} {data.duration_unit}"],
    ], colWidths=[1.55 * inch, 4.75 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#efeee5")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#3d5338")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e7ded1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
    ]))
    story.extend([Spacer(1, 0.15 * inch), table, Spacer(1, 0.25 * inch), Paragraph("Propuesta preparada con la información suministrada por el técnico. Confirma el alcance y las condiciones antes de aceptarla.", styles["Italic"])])
    SimpleDocTemplate(output, pagesize=letter, rightMargin=0.8 * inch, leftMargin=0.8 * inch, topMargin=0.8 * inch, bottomMargin=0.8 * inch).build(story)
    return output.getvalue()


def build_export(data: QuoteDocument, file_format: str) -> tuple[bytes, str, str]:
    filename = safe_filename(data.job_title)
    builders = {
        "txt": (_plain_text, "text/plain; charset=utf-8", "txt"),
        "docx": (_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
        "xlsx": (_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
        "pdf": (_pdf, "application/pdf", "pdf"),
    }
    if file_format in builders:
        builder, content_type, extension = builders[file_format]
        return builder(data), content_type, f"cotizacion-{filename}.{extension}"
    if file_format == "zip":
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(f"cotizacion-{filename}.txt", _plain_text(data))
            archive.writestr(f"cotizacion-{filename}.docx", _docx(data))
            archive.writestr(f"cotizacion-{filename}.xlsx", _xlsx(data))
            archive.writestr(f"cotizacion-{filename}.pdf", _pdf(data))
        return output.getvalue(), "application/zip", f"cotizacion-{filename}.zip"
    raise ValueError("Formato de archivo no compatible.")
